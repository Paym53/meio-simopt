"""
Excel export: writes every table on its own sheet with a one-line
description at the top, bold headers, frozen header row and sensible
column widths.
"""
from __future__ import annotations

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="DDE7F0")

README_ROWS = [
    ("How to read this file", ""),
    ("Week numbering", "Week 1 is the current review week. The model is re-run every week; only week-1 decisions are committed."),
    ("Ages", "Age 1 = arrived this week. A unit of age b has (shelf life - b) weeks left. Units at the oldest "
             "sellable age are scrapped at the end of the week if unsold."),
    ("Order of events in a week",
     "1 ageing -> 2 receipts -> 3a DC order (s,S on DC position) -> 3b production release (capped by RM, "
     "capacity of the production week, batch/MOQ; rest cancelled) -> 3c RM leaves the RMW (reaches "
     "production tau weeks later) -> 3d RM order (s,S on echelon position) -> 4 demand and allocation -> "
     "5 fill recording -> 6 waste and holding cost."),
    ("Policy", "Age-aware (s,S): positions minus expected waste (median lead time window), DC order cap, "
               "minimum physical RM stock at the RMW; the week-1 orders are chosen by lookahead."),
    ("DC inventory position", "FG on hand + FG released but not yet at the DC. The rule uses the effective "
                              "position = inventory position - stock expected to expire before a new order arrives."),
    ("RMW echelon position", "RM on hand + RM on order + BOM x (FG on hand + FG in the pipeline)."),
    ("Allocation", "Oldest age first (FIFO). Within an age: channel with the tightest shelf-life requirement "
                   "first, ties: higher target fill rate. Unmet demand is lost."),
    ("Lead times", "Random and order-preserving: an order never arrives before an earlier order on the same "
                   "lane (arrival = max(order week + lead time, previous arrival)). A release in week t is "
                   "produced in week t + tau (RMW -> PF lead time, deterministic) and reaches the DC in week "
                   "t + tau + L."),
    ("Seeds", "Search, hold-out and test seeds are disjoint random futures (demand + lead times)."),
    ("Fill-rate rules", "Search: mean - Z*SE - margin >= F per cell. Hold-out: weak if mean < F, margin += "
                        "max(bump, F - mean). Final verdict (test seeds): mean >= F, no Z, no margin."),
    ("", ""),
    ("Sheet", "Content"),
    ("01_Settings", "Horizon, evaluation window, seed counts, search settings, run information."),
    ("02_Products", "Finished goods: shelf life, BOM, lot-sizing rules, costs, lead times."),
    ("03_Channels", "Channels: target fill rate, shelf-life requirement, accepted ages, service priority."),
    ("04_Materials", "Raw materials: shelf life, lot-sizing rules, costs, lead times."),
    ("05_Cost_Tiers", "All-units discount tiers for production and PF->DC transport."),
    ("06_Lead_Times", "Lead-time distributions per lane."),
    ("07_Demand_Forecast", "Forecast distribution (mean, sd) per product, channel and week."),
    ("08_Initial_State", "Stock by age and open orders at the start of week 1."),
    ("09_Policy_Schedule", "Week-specific (s,S) levels: start schedule and optimised schedule."),
    ("10_Committed_Week1", "The decisions to execute now: FG order, production release, RM transport, RM orders."),
    ("11_Service_Cells", "Every (week x channel) cell: search, hold-out and test statistics and verdicts."),
    ("12_Holdout_Margins", "Weak cells found by the hold-out check in each round, and the margins applied."),
    ("13_Costs", "Cost per component over the test seeds (mean, sd, P5, P95, SE)."),
    ("14_KPIs", "Pooled fill rates, waste, cancellations, stock levels, order counts (test seeds)."),
    ("15_Weekly_Means", "Mean over all test seeds of every flow, position and cost, per week."),
    ("16_Trace_Weekly", "Trace seeds: every flow, position, fill and cost per week."),
    ("17_Trace_DC_Ages", "Trace seeds: DC stock by age at the start (after receipts) and end (after waste) of each week."),
    ("18_Trace_Sales_Age", "Trace seeds: units sold per channel by age (shows FIFO and shelf-life gates)."),
    ("19_Trace_RM_Ages", "Trace seeds: RMW stock by age at the start and end of each week."),
    ("20_Trace_RM_Shipped", "Trace seeds: RM shipped to production by age (shows FIFO at the RMW)."),
    ("21_Trace_Flows", "Trace seeds: every material flow from location to location (long format, filterable)."),
    ("22_Trace_Orders", "Trace seeds: every order with lead-time draw, planned and actual (order-preserving) arrival."),
    ("23_Checks", "Unit balance checks over all test seeds (must all be 0)."),
    ("24_Search_Log", "Every accepted search step: phase, action, mean cost, feasibility."),
    ("25_Lookahead", "Every week-1 candidate quantity with its mean cost and feasibility."),
]


def _format_sheet(ws, n_header_rows: int) -> None:
    header_row = n_header_rows + 1
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for col_idx, column in enumerate(ws.iter_cols(min_row=header_row, max_row=min(ws.max_row, header_row + 200)), start=1):
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(10, longest + 2), 60)


def write_workbook(path: str, sheets: list[tuple[str, str, pd.DataFrame]], include_readme: bool = True) -> None:
    """sheets: list of (sheet name, one-line description, table)."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        if include_readme:
            _write_readme(writer)
        for name, description, table in sheets:
            table.to_excel(writer, sheet_name=name, index=False, startrow=2)
            ws = writer.sheets[name]
            ws["A1"] = description
            ws["A1"].font = Font(italic=True)
            _format_sheet(ws, n_header_rows=2)


def _write_readme(writer) -> None:
    readme = pd.DataFrame(README_ROWS, columns=["Topic", "Explanation"])
    readme.to_excel(writer, sheet_name="00_ReadMe", index=False)
    ws = writer.sheets["00_ReadMe"]
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 140
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for row in ws.iter_rows(min_row=2):
        if row[0].value in ("How to read this file", "Sheet"):
            row[0].font = Font(bold=True)
            row[1].font = Font(bold=True)
